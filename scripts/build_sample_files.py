"""Render the sample submissions in ``backend/app/sample/documents.py`` to real files.

Developer-only: needs ``reportlab`` and ``Pillow`` (``pip install reportlab pillow``).
The rendered files are committed under ``backend/app/sample/files`` so seeding never
needs these packages. Re-run this script only after editing the sample texts.
"""

from __future__ import annotations

import io
import random
import sys
from html import escape
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "backend"))

from docx import Document  # noqa: E402
from PIL import Image, ImageDraw, ImageFilter, ImageFont  # noqa: E402
from reportlab import rl_config  # noqa: E402
from reportlab.lib import colors  # noqa: E402
from reportlab.lib.enums import TA_CENTER  # noqa: E402
from reportlab.lib.pagesizes import A4  # noqa: E402
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet  # noqa: E402
from reportlab.lib.units import mm  # noqa: E402
from reportlab.lib.utils import ImageReader  # noqa: E402
from reportlab.pdfgen import canvas as pdf_canvas  # noqa: E402
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer  # noqa: E402

from app.sample.documents import essay_documents  # noqa: E402

OUTPUT = ROOT / "backend" / "app" / "sample" / "files"
rl_config.invariant = 1  # reproducible PDF bytes (no creation date or random IDs)


def _styles() -> dict[str, ParagraphStyle]:
    base = getSampleStyleSheet()
    body = ParagraphStyle(
        "Body",
        parent=base["BodyText"],
        fontName="Times-Roman",
        fontSize=10.5,
        leading=14,
        textColor=colors.HexColor("#333333"),
        spaceAfter=8,
    )
    return {
        "title": ParagraphStyle(
            "Title",
            parent=base["Title"],
            fontName="Times-Bold",
            fontSize=19,
            leading=23,
            alignment=TA_CENTER,
            textColor=colors.HexColor("#2F2F2F"),
            spaceAfter=8,
        ),
        "meta": ParagraphStyle(
            "Meta",
            parent=base["Normal"],
            fontName="Helvetica",
            fontSize=8.5,
            leading=11,
            alignment=TA_CENTER,
            textColor=colors.HexColor("#6E655E"),
            spaceAfter=14,
        ),
        "heading": ParagraphStyle(
            "Heading",
            parent=base["Heading1"],
            fontName="Times-Bold",
            fontSize=13,
            leading=16,
            textColor=colors.HexColor("#8F3F24"),
            spaceBefore=4,
            spaceAfter=8,
        ),
        "question": ParagraphStyle(
            "Question",
            parent=body,
            fontName="Times-Italic",
            textColor=colors.HexColor("#555555"),
            spaceAfter=6,
        ),
        "body": body,
        "refs": ParagraphStyle(
            "Refs",
            parent=body,
            fontSize=9.5,
            leading=12.5,
            leftIndent=8 * mm,
            firstLineIndent=-8 * mm,
            spaceAfter=4,
        ),
    }


def _footer(label: str):
    def draw(canvas, doc) -> None:  # type: ignore[no-untyped-def]
        canvas.saveState()
        canvas.setStrokeColor(colors.HexColor("#D8D1C9"))
        canvas.line(22 * mm, 15 * mm, A4[0] - 22 * mm, 15 * mm)
        canvas.setFont("Helvetica", 8)
        canvas.setFillColor(colors.HexColor("#6E655E"))
        canvas.drawString(22 * mm, 10 * mm, label)
        canvas.drawRightString(A4[0] - 22 * mm, 10 * mm, f"Page {doc.page}")
        canvas.restoreState()

    return draw


def _template(target: Path, title: str) -> SimpleDocTemplate:
    return SimpleDocTemplate(
        str(target),
        pagesize=A4,
        leftMargin=22 * mm,
        rightMargin=22 * mm,
        topMargin=20 * mm,
        bottomMargin=22 * mm,
        title=title,
        author="Sample student",
    )


def build_essay(target: Path, spec: dict) -> None:
    styles = _styles()
    story = [Paragraph(escape(spec["title"]), styles["title"]), Paragraph(escape(spec["meta"]), styles["meta"])]
    sections = spec["sections"]
    for index, (heading, paragraphs) in enumerate(sections):
        story.append(Paragraph(escape(heading), styles["heading"]))
        story.extend(Paragraph(escape(text), styles["body"]) for text in paragraphs)
        if spec.get("page_break_each_section") and index < len(sections) - 1:
            story.append(PageBreak())
    story.append(Spacer(1, 6))
    story.append(Paragraph("References", styles["heading"]))
    story.extend(Paragraph(escape(item), styles["refs"]) for item in spec["references"])
    draw = _footer(spec["meta"])
    _template(target, spec["title"]).build(story, onFirstPage=draw, onLaterPages=draw)


def build_quiz(target: Path, spec: dict) -> None:
    styles = _styles()
    story = [Paragraph(escape(spec["title"]), styles["title"]), Paragraph(escape(spec["meta"]), styles["meta"])]
    for number, (_label, question, answer) in enumerate(spec["answers"], 1):
        story.append(Paragraph(f"{number}. Question {number}", styles["heading"]))
        story.append(Paragraph(escape(question), styles["question"]))
        story.append(Spacer(1, 4))
        story.append(Paragraph(escape(answer), styles["body"]))
    draw = _footer(spec["meta"])
    _template(target, spec["title"]).build(story, onFirstPage=draw, onLaterPages=draw)


def build_docx(target: Path, spec: dict) -> None:
    document = Document()
    paragraphs = spec["paragraphs"]
    document.add_heading(paragraphs[0], level=1)
    for text in paragraphs[1:]:
        if text[:2].rstrip(".").isdigit() or text == "References":
            document.add_heading(text, level=2)
        else:
            document.add_paragraph(text)
    core = document.core_properties
    core.author = "Sample student"
    core.title = paragraphs[0]
    buffer = io.BytesIO()
    document.save(buffer)
    target.write_bytes(buffer.getvalue())


def build_scanned(target: Path, spec: dict) -> None:
    """An image-only PDF: no text layer, as produced by a phone scanner."""

    width, height = 1240, 1754  # A4 at 150 dpi
    image = Image.new("L", (width, height), 246)
    draw = ImageDraw.Draw(image)
    try:
        font = ImageFont.load_default(size=30)
    except TypeError:  # Pillow < 10.1
        font = ImageFont.load_default()
    y = 160
    for line in spec["lines"]:
        draw.text((140, y), line, fill=40, font=font)
        y += 52
    rng = random.Random(40231)
    for _ in range(4000):  # scanner speckle
        draw.point((rng.randrange(width), rng.randrange(height)), fill=rng.randrange(170, 230))
    image = image.rotate(0.8, fillcolor=246, resample=Image.BICUBIC).filter(ImageFilter.GaussianBlur(0.6))
    encoded = io.BytesIO()
    image.save(encoded, format="PNG", optimize=True)
    encoded.seek(0)
    canvas = pdf_canvas.Canvas(str(target), pagesize=A4, invariant=1)
    canvas.drawImage(ImageReader(encoded), 0, 0, width=A4[0], height=A4[1])
    canvas.showPage()
    canvas.save()


def main() -> None:
    OUTPUT.mkdir(parents=True, exist_ok=True)
    builders = {
        "essay_pdf": build_essay,
        "quiz_pdf": build_quiz,
        "docx": build_docx,
        "scanned_pdf": build_scanned,
    }
    for name, spec in essay_documents().items():
        builders[spec["kind"]](OUTPUT / name, spec)
        print(f"wrote {name}")


if __name__ == "__main__":
    main()
