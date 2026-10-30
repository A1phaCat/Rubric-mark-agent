"""Lightweight local document extraction for the MVP review flow."""

from __future__ import annotations

import re
from pathlib import Path
from typing import Any

_HEADING = re.compile(r"^(?:\d+(?:\.\d+)*\.?\s+|Appendix\s+[A-Z]\b).{2,100}$", re.IGNORECASE)
_BULLET = re.compile(r"^[•▪◦]\s*")


def extract_document(path: Path, submission_id: str, mime_type: str) -> dict[str, Any]:
    """Extract page-aware paragraphs from a supported uploaded document."""

    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return _extract_pdf(path, submission_id, mime_type)
    if suffix == ".docx":
        return _extract_docx(path, submission_id, mime_type)
    if suffix == ".txt":
        return text_document(submission_id, path.stem, path.read_text(encoding="utf-8", errors="replace"))
    raise ValueError(f"Unsupported submission format: {suffix}")


def text_document(submission_id: str, title: str, text: str, mime_type: str = "text/plain") -> dict[str, Any]:
    """A plain-text submission as one page of paragraphs, the same way an uploaded .txt is read."""

    return _document(submission_id, title, mime_type, [_page(1, text)])


def _extract_pdf(path: Path, submission_id: str, mime_type: str) -> dict[str, Any]:
    from pypdf import PdfReader

    reader = PdfReader(path)
    title = str((reader.metadata or {}).get("/Title") or path.stem)
    pages = [
        _page(number, page.extract_text(extraction_mode="layout") or "")
        for number, page in enumerate(reader.pages, 1)
    ]
    return _document(submission_id, title, mime_type, pages)


def _extract_docx(path: Path, submission_id: str, mime_type: str) -> dict[str, Any]:
    from docx import Document

    source = Document(path)
    lines = [paragraph.text for paragraph in source.paragraphs]
    for table in source.tables:
        lines.extend(" | ".join(cell.text for cell in row.cells) for row in table.rows)
    title = str(source.core_properties.title or path.stem)
    # Word already marks paragraph boundaries; keep each one as its own block
    # instead of re-flowing short paragraphs into their neighbours.
    return _document(submission_id, title, mime_type, [_page(1, "\n\n".join(lines))])


def _page(number: int, text: str) -> dict[str, Any]:
    blocks: list[dict[str, Any]] = []
    current = ""
    paragraph_number = 0

    def flush() -> None:
        nonlocal current, paragraph_number
        value = re.sub(r"\s+", " ", current).strip()
        if value:
            paragraph_number += 1
            blocks.append(
                {
                    "page": number,
                    "label": f"¶{paragraph_number}",
                    "kind": "paragraph",
                    "text": value,
                }
            )
        current = ""

    for raw in text.replace("\r", "").splitlines():
        line = re.sub(r"\s+", " ", raw).strip()
        if not line:
            flush()
            continue
        if _HEADING.match(line) and len(line.split()) <= 14:
            flush()
            blocks.append({"page": number, "label": "", "kind": "heading", "text": line})
            continue
        if _BULLET.match(line):
            flush()
            current = _BULLET.sub("", line)
            if current.endswith((".", "!", "?")):
                flush()
            continue
        if current.endswith("-"):
            current = current[:-1] + line
        else:
            current = f"{current} {line}".strip()
        if len(current) >= 180 and line.endswith((".", "!", "?")):
            flush()
    flush()
    return {"number": number, "blocks": blocks}


def _document(
    submission_id: str, title: str, mime_type: str, pages: list[dict[str, Any]]
) -> dict[str, Any]:
    paragraphs = [block for page in pages for block in page["blocks"]]
    text = " ".join(block["text"] for block in paragraphs)
    return {
        "id": f"doc_{submission_id}",
        "version": 1,
        "page_count": len(pages),
        "word_count": len(re.findall(r"\b\w+\b", text)),
        "content_url": f"/api/v1/submissions/{submission_id}/file",
        "mime_type": mime_type,
        "title": title,
        "paragraphs": paragraphs,
        "extraction": {"method": "text_layer", "status": "complete"},
    }
